"""Text extraction per page (SPEC §10.1).

* PDF → one entry per physical page (pypdf); page numbers are real and shown in citations.
* DOCX → paragraphs, split into ~1200-char sections.
* Markdown / text / CSV → decoded text, split into ~1200-char sections on line boundaries.

Extraction is CPU-bound and synchronous; callers run :func:`extract_text` in a thread.
"""

from __future__ import annotations

import io
import logging
import re

from topskill.knowledge.base import ExtractedText

log = logging.getLogger(__name__)

SECTION_CHARS = 1200
_WS_RE = re.compile(r"[ \t\f\v ]+")
_BLANK_LINES_RE = re.compile(r"\n{3,}")


class UnparseableFile(Exception):
    """The bytes are not a valid document of the declared type."""


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "\n".join(_WS_RE.sub(" ", line).strip() for line in text.split("\n"))
    return _BLANK_LINES_RE.sub("\n\n", text).strip()


def split_sections(text: str, size: int = SECTION_CHARS) -> list[str]:
    """Split text into sections of at most ``size`` chars, preferring line then word boundaries."""
    text = text.strip()
    if not text:
        return []
    sections: list[str] = []
    current = ""
    for line in text.split("\n"):
        while len(line) > size:  # a single very long line: cut at a space when possible
            cut = line.rfind(" ", 0, size)
            cut = cut if cut > size // 2 else size
            head, line = line[:cut].rstrip(), line[cut:].lstrip()
            if current:
                sections.append(current)
                current = ""
            sections.append(head)
        candidate = f"{current}\n{line}" if current else line
        if len(candidate) > size:
            sections.append(current)
            current = line
        else:
            current = candidate
    if current.strip():
        sections.append(current)
    return [s.strip() for s in sections if s.strip()]


def _decode(data: bytes) -> str:
    encoding = "utf-16" if data[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8-sig"
    try:
        return data.decode(encoding)
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _pdf_pages(data: bytes) -> list[str]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise UnparseableFile("encrypted pdf") from exc
        pages: list[str] = []
        for page in reader.pages:
            try:
                pages.append(_tidy(page.extract_text() or ""))
            except Exception:  # one bad page must not lose the others
                log.warning("pdf_page_extract_failed", exc_info=True)
                pages.append("")
        return pages
    except UnparseableFile:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as exc:
        raise UnparseableFile(f"pdf: {type(exc).__name__}") from exc


def _docx_text(data: bytes) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError, OSError) as exc:
        raise UnparseableFile(f"docx: {type(exc).__name__}") from exc
    except Exception as exc:  # zipfile.BadZipFile and lxml errors
        raise UnparseableFile(f"docx: {type(exc).__name__}") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return _tidy("\n".join(parts))


def extract_text(data: bytes, mime: str) -> ExtractedText:
    """Pages of text for an accepted upload. Raises :class:`UnparseableFile` for corrupt files."""
    if mime == "application/pdf":
        return ExtractedText(pages=_pdf_pages(data), real_pages=True)
    if mime.endswith("wordprocessingml.document"):
        return ExtractedText(pages=split_sections(_docx_text(data)), real_pages=False)
    return ExtractedText(pages=split_sections(_tidy(_decode(data))), real_pages=False)
